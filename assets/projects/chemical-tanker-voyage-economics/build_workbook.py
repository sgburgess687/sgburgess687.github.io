"""
build_workbook.py — generates model/chemical_tanker_voyage_model.xlsx
All model cells are live Excel formulas (no hardcoded results). Assumption values are
imported from analysis.py so the workbook and the Python engine share one source of truth.
Independent analysis using publicly available data. Not affiliated with or endorsed by any organization named.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import analysis as AN

HERE = Path(__file__).resolve().parent
OUT = HERE / "model" / "chemical_tanker_voyage_model.xlsx"
A = AN.ASSUMPTIONS

FONT = "Arial"
NAVY = "1F3A5F"
BLUE_IN = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
H1 = Font(name=FONT, bold=True, size=16, color=NAVY)
H2 = Font(name=FONT, bold=True, size=12, color=NAVY)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
SWITCH_FILL = PatternFill("solid", fgColor="FFFF00")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
OUT_FILL = PatternFill("solid", fgColor="E8EEF5")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
DISCLAIMER = ("Independent analysis using publicly available data. Not affiliated with or endorsed by any "
              "organization named.")

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'

wb = Workbook()
names = {}


def name(nm, ws, cell):
    ref = f"'{ws.title}'!${''.join(c for c in cell if c.isalpha())}${''.join(c for c in cell if c.isdigit())}"
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    names[nm] = ref


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, italic=f.italic, size=f.size, color=f.color)


def header(ws, row, cols, start=1):
    for i, h in enumerate(cols):
        c = ws.cell(row=row, column=start + i, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = H1
    ws["A2"] = DISCLAIMER
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="7F7F7F")
    if sub:
        ws["A3"] = sub
        ws["A3"].font = Font(name=FONT, size=10, color="404040")
    ws.sheet_view.showGridLines = False


# =============================================================================
# Sheet order
# =============================================================================
cover = wb.active
cover.title = "Cover"
inp = wb.create_sheet("Inputs")
rt = wb.create_sheet("Routes")
pt = wb.create_sheet("Port_Tariffs")
cn = wb.create_sheet("Canal_Tolls")
pnl = wb.create_sheet("Voyage_PnL")
sens = wb.create_sheet("Sensitivity")
be = wb.create_sheet("Breakeven")
eng = wb.create_sheet("Engine")
db = wb.create_sheet("Data_Brent")
dp = wb.create_sheet("Data_PPI")

# =============================================================================
# Data sheets
# =============================================================================
brent, ppi = AN.load_data()
title(db, "Brent crude spot price, monthly average (US$/bbl)",
      "Source: FRED series MCOILBRENTEU (EIA, Europe Brent Spot Price FOB), https://fred.stlouisfed.org/series/MCOILBRENTEU, "
      "retrieved 17-Sep-2026. Used as the bunker-price driver.")
header(db, 5, ["Month", "Brent $/bbl"])
B0 = 6
for i, (d, v) in enumerate(brent.items()):
    db.cell(row=B0 + i, column=1, value=d.to_pydatetime()).number_format = "mmm-yyyy"
    db.cell(row=B0 + i, column=2, value=float(v)).number_format = NUM2
nB = len(brent)
BR = f"$B${B0}:$B${B0 + 499}"
db["D5"], db["E5"] = "Statistic", "Value"
header(db, 5, ["Statistic", "Value", "Definition"], start=4)
stats = [
    ("Count of months", f"=COUNT({BR})", "", NUM),
    ("Latest month value", f"=INDEX({BR},E6)", "Last observation", NUM2),
    ("Trailing 12-month average", f"=ROUND(AVERAGE(INDEX({BR},E6-11):INDEX({BR},E6)),2)", "Base-case bunker driver", NUM2),
    ("5-year average", f"=ROUND(AVERAGE(INDEX({BR},E6-59):INDEX({BR},E6)),2)", "Last 60 months", NUM2),
    ("5-year P10", f"=ROUND(PERCENTILE(INDEX({BR},E6-59):INDEX({BR},E6),0.1),2)", "Tornado low", NUM2),
    ("5-year P90", f"=ROUND(PERCENTILE(INDEX({BR},E6-59):INDEX({BR},E6),0.9),2)", "Tornado high", NUM2),
    ("5-year minimum", f"=MIN(INDEX({BR},E6-59):INDEX({BR},E6))", "", NUM2),
    ("5-year maximum", f"=MAX(INDEX({BR},E6-59):INDEX({BR},E6))", "", NUM2),
]
for i, (lab, f, d, fmt) in enumerate(stats):
    r = 6 + i
    db.cell(row=r, column=4, value=lab)
    c = db.cell(row=r, column=5, value=f)
    c.number_format = fmt
    db.cell(row=r, column=6, value=d)
name("Brent_Latest", db, "E7"); name("Brent_TTM", db, "E8"); name("Brent_5Y", db, "E9")
name("Brent_P10", db, "E10"); name("Brent_P90", db, "E11")
db["D16"] = ("Note: values transcribed from the FRED text file via a web tool on 17-Sep-2026; "
             "re-download the CSV to confirm before publishing.")
db["D16"].font = Font(name=FONT, italic=True, color="C00000")
db.column_dimensions["A"].width = 12; db.column_dimensions["B"].width = 13
db.column_dimensions["D"].width = 28; db.column_dimensions["E"].width = 12; db.column_dimensions["F"].width = 26
ch = LineChart(); ch.title = "Brent monthly average ($/bbl)"; ch.height = 8; ch.width = 18
ch.add_data(Reference(db, min_col=2, min_row=5, max_row=B0 + nB - 1), titles_from_data=True)
ch.set_categories(Reference(db, min_col=1, min_row=B0, max_row=B0 + nB - 1)); ch.legend = None
ch.x_axis.number_format = "yyyy"; ch.x_axis.majorTimeUnit = "years"
db.add_chart(ch, "H5")

title(dp, "BLS Producer Price Index — Deep Sea Freight Transportation (Dec-2008 = 100)",
      "Source: BLS series WPU30130101 via FRED, https://fred.stlouisfed.org/series/WPU30130101, file supplied in "
      "tanker_voyage_data/freight_rates (retrieved Sep-2026). A broad freight proxy, NOT a chemical-tanker rate.")
header(dp, 5, ["Month", "Index", "YoY change"])
P0 = 6
first2016 = None
for i, (d, v) in enumerate(ppi.items()):
    r = P0 + i
    dp.cell(row=r, column=1, value=d.to_pydatetime()).number_format = "mmm-yyyy"
    dp.cell(row=r, column=2, value=float(v)).number_format = NUM1
    if i >= 12:
        dp.cell(row=r, column=3, value=f"=B{r}/B{r-12}-1").number_format = PCT
    if first2016 is None and d.year >= 2016:
        first2016 = r
lastP = P0 + len(ppi) - 1
header(dp, 5, ["Statistic", "Value", "Definition"], start=5)
pst = [
    ("Latest index value", f"=B{lastP}", "Jun-2026", NUM1),
    ("YoY change P10 (2016+)", f"=ROUND(PERCENTILE(C{first2016}:C{lastP},0.1),3)", "Weak-market freight band", PCT),
    ("YoY change P90 (2016+)", f"=ROUND(PERCENTILE(C{first2016}:C{lastP},0.9),3)", "Strong-market freight band", PCT),
    ("Freight lever low (x)", "=1+ROUND(F7,2)", "Tornado low multiplier", MULT),
    ("Freight lever high (x)", "=1+ROUND(F8,2)", "Tornado high multiplier", MULT),
]
for i, (lab, f, d, fmt) in enumerate(pst):
    dp.cell(row=6 + i, column=5, value=lab)
    c = dp.cell(row=6 + i, column=6, value=f); c.number_format = fmt
    dp.cell(row=6 + i, column=7, value=d)
name("Freight_Low", dp, "F9"); name("Freight_High", dp, "F10")
dp["E13"] = ("Why this series: public chemical-tanker rate assessments are subscription-only. The PPI's year-on-year "
             "distribution is used only to size a plausible freight-rate band; spot chemical rates are more volatile, "
             "so this band is likely conservative.")
dp["E13"].alignment = Alignment(wrap_text=True, vertical="top"); dp.merge_cells("E13:G17")
dp.column_dimensions["A"].width = 12; dp.column_dimensions["E"].width = 26; dp.column_dimensions["F"].width = 11
dp.column_dimensions["G"].width = 28
ch = LineChart(); ch.title = "BLS PPI deep-sea freight (index)"; ch.height = 8; ch.width = 18
ch.add_data(Reference(dp, min_col=2, min_row=5, max_row=lastP), titles_from_data=True)
ch.set_categories(Reference(dp, min_col=1, min_row=P0, max_row=lastP)); ch.legend = None
ch.x_axis.number_format = "yyyy"
dp.add_chart(ch, "I5")

# =============================================================================
# Inputs
# =============================================================================
title(inp, "Inputs & scenario switches",
      "Yellow cells = scenario switches. Blue = hardcoded assumption (edit freely). Black = formula. "
      "Every assumption carries its basis in column D.")
for col, w in zip("ABCDE", (38, 14, 16, 70, 12)):
    inp.column_dimensions[col].width = w

r = 5
inp.cell(row=r, column=1, value="SCENARIO SWITCHES").font = H2
r += 1
header(inp, r, ["Switch", "Selection", "Resolves to", "Options"])
dv = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
dv6 = DataValidation(type="list", formula1='"1,2,3,4,5,6"', allow_blank=False)
inp.add_data_validation(dv); inp.add_data_validation(dv6)

switch_rows = {}


def switch(label, key, val, resolves, options, v=dv):
    global r
    r += 1
    inp.cell(row=r, column=1, value=label).font = BOLD
    c = inp.cell(row=r, column=2, value=val); c.font = BLUE_IN; c.fill = SWITCH_FILL; c.border = BOX
    c.alignment = Alignment(horizontal="center")
    v.add(c)
    inp.cell(row=r, column=3, value=resolves)
    inp.cell(row=r, column=4, value=options).alignment = Alignment(wrap_text=True)
    name(key, inp, f"B{r}")
    switch_rows[key] = r


switch("1. Route", "Sw_Route", 1, "=INDEX(Routes!$B$6:$B$9,Sw_Route)",
       "1 USG-ARA transatlantic | 2 USG-Far East via Panama | 3 ARA-India WC via Suez | 4 ARA-India WC via Cape")
switch("2. Bunker price case", "Sw_Bunker", 1, None,
       "1 Trailing-12m avg | 2 Latest month | 3 5-yr avg | 4 5-yr P10 | 5 5-yr P90 | 6 Custom (below)", v=dv6)
switch("3. Freight market case", "Sw_Freight", 1, None,
       "1 Base rates | 2 Weak (PPI YoY P10) | 3 Strong (PPI YoY P90) | 4 Custom multiplier (below)")
switch("4. Port congestion case", "Sw_Port", 1, None,
       "1 Normal (1.0x port days) | 2 Congested (1.5x) | 3 Severe (2.0x) | 4 Custom (below)")
switch("5. Cargo utilization case", "Sw_Util", 1, None,
       "1 Base (85% / 50%) | 2 Soft (70% / 25%) | 3 Firm (95% / 75%) | 4 Custom (below)  [fronthaul / backhaul]")

r += 2
inp.cell(row=r, column=1, value="CUSTOM OVERRIDES (used only when a switch selects 'Custom')").font = H2
r += 1
header(inp, r, ["Item", "Value", "Unit", "Note"])


def row_in(label, key, val, unit, basis, fmt=None, key_assump=False):
    global r
    r += 1
    inp.cell(row=r, column=1, value=label)
    c = inp.cell(row=r, column=2, value=val)
    c.font = BLUE_IN if not (isinstance(val, str) and val.startswith("=")) else BLACK
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if key_assump:
        c.fill = KEY_FILL
    inp.cell(row=r, column=3, value=unit)
    inp.cell(row=r, column=4, value=basis).alignment = Alignment(wrap_text=True, vertical="top")
    if key:
        name(key, inp, f"B{r}")


row_in("Custom Brent", "Cust_Brent", 80.0, "$/bbl", "User entry", NUM2)
row_in("Custom freight multiplier", "Cust_Freight", 1.0, "x base rate", "User entry", MULT)
row_in("Custom port-time multiplier", "Cust_Port", 1.0, "x base port days", "User entry", MULT)
row_in("Custom fronthaul utilization", "Cust_UtilFH", 0.85, "% capacity", "User entry", PCT)
row_in("Custom backhaul utilization", "Cust_UtilBH", 0.50, "% capacity", "User entry (0% = ballast return)", PCT)

r += 2
inp.cell(row=r, column=1, value="ACTIVE LEVERS (what the P&L uses — formulas, do not overwrite)").font = H2
r += 1
header(inp, r, ["Lever", "Active value", "Unit", "How resolved"])
row_in("Brent crude", "Brent", "=CHOOSE(Sw_Bunker,Brent_TTM,Brent_Latest,Brent_5Y,Brent_P10,Brent_P90,Cust_Brent)",
       "$/bbl", "Switch 2 -> Data_Brent statistics (FRED MCOILBRENTEU)", NUM2)
row_in("Freight rate multiplier", "FreightMult", "=CHOOSE(Sw_Freight,1,Freight_Low,Freight_High,Cust_Freight)",
       "x base rate", "Switch 3 -> Data_PPI percentile bands", MULT)
row_in("Port-time multiplier", "PortMult", "=CHOOSE(Sw_Port,1,1.5,2,Cust_Port)", "x base port days",
       "Switch 4 (scenario bands, not a pulled series — USACE LPMS delay microdata is not public)", MULT)
row_in("Fronthaul utilization", "UtilFH", "=CHOOSE(Sw_Util,0.85,0.7,0.95,Cust_UtilFH)", "% capacity", "Switch 5", PCT)
row_in("Backhaul utilization", "UtilBH", "=CHOOSE(Sw_Util,0.5,0.25,0.75,Cust_UtilBH)", "% capacity", "Switch 5", PCT)
inp.cell(row=switch_rows["Sw_Bunker"], column=3, value="=TEXT(Brent,\"$0.00\")&\"/bbl Brent\"")
inp.cell(row=switch_rows["Sw_Freight"], column=3, value="=TEXT(FreightMult,\"0.00\")&\"x base rates\"")
inp.cell(row=switch_rows["Sw_Port"], column=3, value="=TEXT(PortMult,\"0.00\")&\"x port days\"")
inp.cell(row=switch_rows["Sw_Util"], column=3, value="=TEXT(UtilFH,\"0%\")&\" / \"&TEXT(UtilBH,\"0%\")")

r += 2
inp.cell(row=r, column=1, value="VESSEL — generic 38,000 dwt IMO II stainless-steel parcel chemical tanker").font = H2
r += 1
header(inp, r, ["Assumption", "Value", "Unit", "Basis / judgment (ESTIMATE unless a source is named)"])
row_in("Deadweight", "DWT", A["dwt"], "t", "Typical deep-sea parcel tanker size class (32-45k dwt). Generic, not a specific ship.", NUM, True)
row_in("Usable cargo capacity", "Capacity", A["cargo_capacity_t"], "t", "DWT less ~2,000 t bunkers, fresh water, stores, constant.", NUM, True)
row_in("Operating speed", "Speed", A["speed_kn"], "knots", "Lever. Typical laden service speed 13-14.5 kn for this class.", NUM1, True)
row_in("Reference speed for consumption", "RefSpeed", A["ref_speed_kn"], "knots", "Speed at which the consumption figures below apply.", NUM1)
row_in("Sea consumption, laden", "ConsLaden", A["cons_laden_tpd"], "t/day", "At reference speed. Non-eco 2010s-built design; eco ships ~20-24 t/day.", NUM1, True)
row_in("Sea consumption, ballast", "ConsBallast", A["cons_ballast_tpd"], "t/day", "At reference speed; applies to backhaul leg only when backhaul utilization = 0%.", NUM1)
row_in("Port / canal consumption (MGO)", "ConsPort", A["cons_port_tpd"], "t/day", "Cargo pumps, heating, inerting, hotel load. Canal transit days burn at this rate (simplification).", NUM1, True)
row_in("Consumption-speed exponent", "SpeedExp", A["speed_exponent"], "", "Admiralty (cube) law; fuel per mile scales with speed squared.", NUM1)
row_in("Sea margin", "SeaMargin", A["sea_margin"], "% of distance", "Weather and routeing allowance.", PCT)
row_in("Panama Canal net tonnage (PC/UMS)", "PCUMS", A["pc_ums_t"], "t", "Estimate for class; use the ship's certificate in practice.", NUM)
row_in("Suez Canal net tonnage (SCNT)", "SCNT", A["scnt_t"], "t", "Estimate for class; use the ship's certificate in practice.", NUM)
row_in("Hull & machinery value", "HullValue", A["hull_value_usd"], "$", "Used only for war-risk additional premium. Estimate.", USD)

r += 2
inp.cell(row=r, column=1, value="FUEL, CARBON & COMMERCIAL").font = H2
r += 1
header(inp, r, ["Assumption", "Value", "Unit", "Basis / judgment"])
row_in("VLSFO price per $1/bbl Brent", "VLSFO_Mult", A["vlsfo_per_brent"], "$/t per $/bbl",
       "ESTIMATE. ~7.3 bbl/t conversion x ~5% product premium. At $82 Brent -> ~$631/t VLSFO. Spot-check vs Ship & Bunker Houston/Rotterdam.", NUM2, True)
row_in("MGO price per $1/bbl Brent", "MGO_Mult", A["mgo_per_brent"], "$/t per $/bbl",
       "ESTIMATE. Distillate premium over crude. At $82 Brent -> ~$795/t MGO. Used in ECAs and in port.", NUM2, True)
row_in("Freight commission (address + brokerage)", "Commission", A["commission"], "% gross freight", "Market-standard 1.25% + 1.25%.", PCT)
row_in("Misc. voyage costs per round voyage", "MiscVoy", A["misc_voyage_usd"], "$", "Tank cleaning, surveys, nitrogen, sundries. ESTIMATE.", USD)
row_in("Port disbursement multiplier", "PortCostMult", A["port_cost_mult"], "x", "Lever on all per-call port costs (Routes sheet).", MULT)
row_in("Daily vessel cost benchmark", "Benchmark", A["daily_cost_benchmark"], "$/day",
       "ESTIMATE. Owner: OPEX ~$8.5k + capital recovery ~$9.5k. Operator: substitute your T/C-in hire. TCE above this = voyage adds value.", USD, True)
row_in("EU ETS allowance price", "EUA", A["eua_eur"], "EUR/t CO2", "ESTIMATE — check ICE EUA futures on the day.", NUM2, True)
row_in("USD per EUR", "FX", A["usd_per_eur"], "$/EUR", "ESTIMATE.", NUM2)
row_in("ETS phase-in (2026)", "PhaseIn", A["ets_phase_in"], "%", "EU ETS maritime: 40% (2024), 70% (2025), 100% from 2026 emissions.", PCT)
row_in("CO2 factor, VLSFO", "CO2_VLSFO", A["co2_vlsfo"], "t CO2 / t fuel", "EU MRV / IMO emission factor for HFO-type fuels.", '0.000')
row_in("CO2 factor, MGO", "CO2_MGO", A["co2_mgo"], "t CO2 / t fuel", "EU MRV / IMO emission factor for MDO/MGO.", '0.000')

r += 2
inp.cell(row=r, column=1, value="VESSEL TONNAGE & DIMENSIONS USED BY PORT TARIFFS").font = H2
r += 1
header(inp, r, ["Assumption", "Value", "Unit", "Basis / judgment"])
row_in("Gross tonnage (GT)", "GT", A["gt"], "GT",
       "Estimate for class. Port of Rotterdam's own worked example uses a 23,230 GT chemical tanker, which corroborates the order of magnitude.", NUM, True)
row_in("Length overall", "LOA_ft", A["loa_ft"], "feet",
       "~183 m. Sets the Port Houston dockage band (600-650 ft).", NUM)

r += 2
inp.cell(row=r, column=1, value="CANAL-RELATED").font = H2
r += 1
header(inp, r, ["Assumption", "Value", "Unit", "Basis / judgment"])
row_in("USD per SDR", "SDR_USD", A["sdr_usd"], "$/SDR", "ESTIMATE — check IMF daily rate (imf.org rms_five).", '0.0000')
row_in("Suez chemical-tanker surcharge, laden", "SuezSurL", A["suez_surcharge_laden"], "%",
       "SCA Circular 7/2023 as reported (20% for chemical tankers). Applied on petroleum-product-tanker bands as proxy.", PCT)
row_in("Suez surcharge, ballast", "SuezSurB", A["suez_surcharge_ballast"], "%", "SCA Circular 7/2023 as reported.", PCT)
row_in("Suez ancillary costs per transit", "SuezAnc", A["suez_ancillary_usd"], "$", "Pilotage, tugs, mooring, agency. ESTIMATE.", USD)
row_in("War-risk additional premium per Suez transit", "WarAP", A["war_risk_ap_pct"], "% hull value",
       "ILLUSTRATIVE. Red Sea AP has swung widely since 2023; set to 0 to test a normalized Suez.", '0.00%', True)
row_in("Panama linehandlers per transit", "Linehandlers", A["panama_linehandlers"], "count", "ESTIMATE.", NUM)
inp.freeze_panes = "A5"

# =============================================================================
# Routes & regions
# =============================================================================
title(rt, "Routes and port regions",
      "Round voyage: fronthaul A->B laden, backhaul B->A. Distances APPROXIMATE (verify with a distance table). "
      "Freight rates are CALIBRATED ASSUMPTIONS, not observed fixtures. Port costs come from the Port_Tariffs sheet.")
header(rt, 5, ["#", "Route", "Region A", "Region B", "Distance per leg (nm)", "ECA miles per leg (nm)", "Canal",
               "Canal days per transit", "EU ETS share of sea emissions", "Base freight fronthaul ($/t)",
               "Base freight backhaul ($/t)", "Port days per op, A", "Port days per op, B",
               "Fixed port cost per op, A ($)", "Fixed port cost per op, B ($)", "EU port A (1/0)", "EU port B (1/0)",
               "Canal cost per laden transit ($)", "Canal cost per ballast transit ($)",
               "Cargo-based port rate A ($/t)", "Cargo-based port rate B ($/t)",
               "Cargo-component cap A ($)", "Cargo-component cap B ($)",
               "Berth-day share A", "Berth-day share B"])
REG_ROW0 = 14
for i, (rid, ro) in enumerate(AN.ROUTES.items()):
    rr = 6 + i
    vals = [rid, ro["name"], ro["A"], ro["B"], ro["nm"], ro["eca_nm"], ro["canal"], ro["canal_days"],
            ro["eu_share"], ro["rate_fh"], ro["rate_bh"]]
    fmts = [None, None, None, None, NUM, NUM, None, NUM1, PCT, USD2, USD2]
    for j, (v, f) in enumerate(zip(vals, fmts)):
        c = rt.cell(row=rr, column=1 + j, value=v)
        c.font = BLUE_IN if j >= 2 else BLACK
        if f:
            c.number_format = f
        c.border = BOX
    reg = f"$A${REG_ROW0 + 1}:$A${REG_ROW0 + 4}"
    # (target column, region-table source column, number format)
    lookups = [(12, "C", NUM1), (13, "C", NUM1), (14, "D", USD), (15, "D", USD), (16, "H", None), (17, "H", None),
               (20, "E", USD2), (21, "E", USD2), (22, "F", USD), (23, "F", USD), (24, "G", PCT), (25, "G", PCT)]
    for col, src, fmt in lookups:
        key = "C" if col % 2 == 0 else "D"  # even target cols read region A, odd read region B
        c = rt.cell(row=rr, column=col,
                    value=f"=INDEX(${src}${REG_ROW0+1}:${src}${REG_ROW0+4},MATCH({key}{rr},{reg},0))")
        if fmt:
            c.number_format = fmt
        c.border = BOX
    rt.cell(row=rr, column=18, value=f'=IF(G{rr}="Panama",Canal_Tolls!$C$20,IF(G{rr}="Suez",Canal_Tolls!$C$45,0))').number_format = USD
    rt.cell(row=rr, column=19, value=f'=IF(G{rr}="Panama",Canal_Tolls!$D$20,IF(G{rr}="Suez",Canal_Tolls!$D$45,0))').number_format = USD
    for j in (18, 19):
        rt.cell(row=rr, column=j).border = BOX
dvc = DataValidation(type="list", formula1='"None,Panama,Suez"'); rt.add_data_validation(dvc); dvc.add("G6:G9")
rt.cell(row=10, column=2, value="Distances: approximate great-circle/standard routeing figures (Houston-Rotterdam ~5,000 nm; "
        "Houston-Ulsan via Panama ~9,650 nm; Rotterdam-Mumbai ~6,350 nm via Suez / ~10,850 nm via Cape). "
        "Verify with a distance table (e.g., Netpas, BP Shipping Marine Distances).").font = Font(name=FONT, italic=True, size=9)

rt.cell(row=REG_ROW0 - 1, column=1, value="PORT REGIONS — port cost per operation = fixed + MIN(cargo tonnes x cargo rate, cap) + dockage (Port_Tariffs sheet)").font = H2
header(rt, REG_ROW0, ["Code", "Region", "Port days per operation (incl. waiting)", "Fixed cost per op ($)",
                      "Cargo rate ($/t)", "Cargo cap ($)", "Berth-day share", "EU port (1/0)", "Basis"])
region_cells = {
    "USG": ("=Port_Tariffs!$C$14", "0", "0", "=Port_Tariffs!$C$11",
            "PUBLISHED: Port Houston Tariff 14 harbor fee + LOA-based dockage (see Port_Tariffs). Pilotage, tugs, linesmen, "
            "agency and private-terminal charges are an estimate."),
    "ARA": ("=Port_Tariffs!$C$29", "=Port_Tariffs!$C$30", "=Port_Tariffs!$C$31", "0",
            "PUBLISHED: Port of Rotterdam seaport dues — GT component + cargo component (capped) + sustainability component, "
            "2025 schedule indexed 3.5% for 2026. Pilotage, tugs, boatmen and agency are an estimate."),
    "FEA": ("65000", "0", "0", "0", "ESTIMATE for Korea/China discharge — no published schedule transcribed."),
    "IND": ("55000", "0", "0", "0", "ESTIMATE for Kandla/Mumbai/JNPT discharge — no published schedule transcribed."),
}
for i, (code, g) in enumerate(AN.REGIONS.items()):
    rr = REG_ROW0 + 1 + i
    fixed, rate, cap, berth, basis_txt = region_cells[code]
    vals = [code, g["name"], g["days"], fixed, rate, cap, berth, g["eu"], basis_txt]
    fmts = [None, None, NUM1, USD, USD2, USD, PCT, None, None]
    for j, (v, f) in enumerate(zip(vals, fmts)):
        c = rt.cell(row=rr, column=1 + j, value=v)
        c.font = GREEN if isinstance(v, str) and v.startswith("=") else (BLUE_IN if j in (2, 7) else BLACK)
        if f:
            c.number_format = f
        c.border = BOX
    rt.cell(row=rr, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    rt.row_dimensions[rr].height = 46
    rt.merge_cells(start_row=rr, start_column=9, end_row=rr, end_column=16)
rt.row_dimensions[5].height = 58
for col, w in zip(range(1, 26), (5, 26, 9, 9, 11, 11, 9, 10, 11, 12, 12, 11, 11, 13, 13, 9, 9, 14, 14, 11, 11, 12, 12, 10, 10)):
    rt.column_dimensions[L(col)].width = w

# =============================================================================
# Port tariffs
# =============================================================================
title(pt, "Port cost build-up from published tariffs",
      "Blue cells are published tariff rates (sources named in column E) or estimated ancillary costs. "
      "Only the US Gulf and ARA legs have transcribed published schedules; Far East and India are lump-sum estimates.")
pt.column_dimensions["A"].width = 48
for col in "BCD":
    pt.column_dimensions[col].width = 15
pt.column_dimensions["E"].width = 72



def dock_formula(d):
    """Port Houston Tariff 14 Subrule 095 taper, as an Excel formula for a days expression."""
    return (f"IF({d}<=2,{d},IF({d}<=3,2+0.9*({d}-2),IF({d}<=4,2.9+0.75*({d}-3),"
            f"IF({d}<=5,3.65+0.6*({d}-4),4.25+0.5*({d}-5)))))")


def pt_row(rr, label, value, fmt, note, calc=False):
    pt.cell(row=rr, column=1, value=label)
    c = pt.cell(row=rr, column=3, value=value)
    c.number_format = fmt
    c.font = BLACK if calc else BLUE_IN
    pt.cell(row=rr, column=5, value=note).alignment = Alignment(wrap_text=True, vertical="top")
    return c


pt["A5"] = "PORT HOUSTON — Tariff No. 14, effective 1 January 2025"; pt["A5"].font = H2
pt["A6"] = "Source: porthouston.com/toolbox/rates/tariffs (Tariff 14 PDF), Subrules 095 (dockage) and 096 (harbor fee)"
pt["A6"].font = Font(name=FONT, italic=True, size=9, color="595959")
header(pt, 7, ["Item", "", "Value", "", "Source / note"])
pt_row(8, "Harbor fee, vessels 250 ft and over", A["ph_harbor_fee"], USD2, "Subrule 096, per entry into the Port Authority's jurisdictional limits. Published rate.")
pt_row(9, "Dockage rate, LOA 600-650 ft band", A["ph_dockage_per_ft"], USD2, "Subrule 095, $ per foot of LOA per 24 hours. Published rate. Applies at Port Authority wharves; private liquid terminals set their own dockage.")
pt_row(10, "Length overall (feet)", "=LOA_ft", NUM, "From Inputs.", calc=True)
pt_row(11, "Share of port days alongside (rest is waiting)", A["usg_berth_share"], PCT, "ESTIMATE. Waiting at anchorage incurs no dockage, so only berth time is charged.")
pt_row(12, "Berth days per operation, at the active port-time case", "=Routes!$C$15*C11*PortMult", NUM2, "Region port days x berth share x port-time multiplier.", calc=True)
pt_row(13, "Dockage day-factor (tariff taper)", "=" + dock_formula("C12"), NUM2, "Subrule 095: days 1-2 at full rate, day 3 at 90%, day 4 at 75%, day 5 at 60%, day 6+ at 50%.", calc=True)
pt_row(14, "FIXED cost per operation (harbor fee + ancillary)", "=C8+C15", USD, "Dockage is added per operation inside the model, because it moves with the congestion case.", calc=True)
pt_row(15, "Ancillary: pilotage in/out, tugs, linesmen, agency, terminal", A["ph_ancillary"], USD, "ESTIMATE — replace with an agency proforma disbursement account.")
pt_row(16, "Illustrative dockage at the active case", "=LOA_ft*C9*C13", USD, "For reference; the model computes this per operation.", calc=True)

pt["A20"] = "PORT OF ROTTERDAM (ARA) — seaport dues, 2025 schedule indexed for 2026"; pt["A20"].font = H2
pt["A21"] = ("Source: Port of Rotterdam 'General terms and conditions including port tariffs', Annex 1 Tables 1-3 and "
             "the 133.3% chemical-tanker efficiency cap; portofrotterdam.com/en/port-dues-tariffs")
pt["A21"].font = Font(name=FONT, italic=True, size=9, color="595959")
header(pt, 22, ["Item", "", "Value", "", "Source / note"])
pt_row(23, "GT tariff, chemical/gas tanker (EUR per GT)", A["rtm_gt_rate"], '0.000', "Table 1, rate type C. Published rate.")
pt_row(24, "Sustainability component (EUR per GT)", A["rtm_sust_rate"], '0.000', "Table 3. No ESI or Green Award discount assumed — a Green Award certificate would cut this component by 70%.")
pt_row(25, "Cargo rate, other liquid bulk (EUR per tonne)", A["rtm_cargo_rate"], '0.000', "Table 2, commodity type 07. Published rate.")
pt_row(26, "Cargo-component cap, chemical tankers", A["rtm_cap_pct"], PCT, "Efficiency discount: cargo dues capped at GT x 133.3% x cargo rate. Published rule.")
pt_row(27, "2026 indexation", A["rtm_index"], '0.000', "Port of Rotterdam published a 3.5% indexation of port tariffs for 2026; applied to the 2025 schedule.")
pt_row(28, "Ancillary: pilotage, tugs, boatmen, agency, terminal", A["rtm_ancillary"], USD, "ESTIMATE — replace with an agency proforma disbursement account.")
pt_row(29, "FIXED cost per operation ($)", "=GT*(C23+C24)*C27*FX+C28", USD, "Vessel and sustainability components, converted at the USD/EUR rate on Inputs.", calc=True)
pt_row(30, "Cargo rate per tonne ($)", "=C25*C27*FX", USD2, "Applied to the tonnes loaded or discharged at this port.", calc=True)
pt_row(31, "Cargo-component cap ($)", "=GT*C26*C25*C27*FX", USD, "", calc=True)
pt_row(32, "Illustrative dues, 30,600 t discharge", "=C29-C28+MIN(30600*C30,C31)", USD, "Sanity check against the worked example in the port's own annex.", calc=True)

pt["A36"] = "FAR EAST AND INDIA — no published schedule transcribed; lump-sum estimates on the Routes sheet."
pt["A36"].font = Font(name=FONT, italic=True, size=10, color="C00000")
for rr in (14, 29, 30, 31):
    pt.cell(row=rr, column=3).fill = OUT_FILL
    pt.cell(row=rr, column=1).font = BOLD

# =============================================================================
# Canal tolls
# =============================================================================
title(cn, "Canal toll build-up (per transit)",
      "Tariff values in blue are taken from the supplied extracts of official schedules (see data/README.md). "
      "Vessel tonnages and ancillaries are estimates.")
cn.column_dimensions["A"].width = 44
for col in "BCDE":
    cn.column_dimensions[col].width = 16
cn.column_dimensions["F"].width = 60
cn["A5"] = "PANAMA CANAL — 'Super' vessel (beam > 91 ft), chemical tanker"; cn["A5"].font = H2
header(cn, 6, ["Item", "Tariff", "Laden ($)", "Ballast ($)", "", "Source / note"])
pan = [
    ("Fixed tariff per transit (1010.FS01)", 100000, "=B7", "=B7*$B$18", "ACP tariff schedule; ballast pays 85% (1010.BA01)"),
    ("Capacity tariff per PC/UMS t, chemical (1010.QS01)", 4.00, "=B8*PCUMS", "=B8*PCUMS*$B$18", "ACP tariff schedule x PC/UMS (Inputs)"),
    ("Mandatory tug, complete transit (Panamax locks)", 7000, "=B9", "=B9", "ACP schedule (regular-vessel rate used as proxy)"),
    ("Linehandling per linehandler (Panamax locks)", 270, "=B10*Linehandlers", "=B10*Linehandlers", "ACP schedule x linehandler count (Inputs)"),
    ("Transit reservation (Panamax)", 10500, "=B11", "=B11", "ACP schedule; optional in practice, included as base"),
    ("Security charge (PC/UMS >= 3,000)", 1250, "=B12", "=B12", "ACP schedule"),
    ("Fresh water surcharge, fixed (LOA > 91.44 m)", 10000, "=B13", "=B13", "ACP schedule; variable portion omitted"),
    ("Port pilotage, per call (x2: each end)", 1275, "=B14*2", "=B14*2", "ACP schedule"),
]
for i, (lab, t, fl, fb, note) in enumerate(pan):
    rr = 7 + i
    cn.cell(row=rr, column=1, value=lab)
    c = cn.cell(row=rr, column=2, value=t); c.font = BLUE_IN; c.number_format = USD2
    cn.cell(row=rr, column=3, value=fl).number_format = USD
    cn.cell(row=rr, column=4, value=fb).number_format = USD
    cn.cell(row=rr, column=6, value=note)
cn["A18"] = "Ballast factor (1010.BA01)"; cn["B18"] = 0.85; cn["B18"].font = BLUE_IN; cn["B18"].number_format = PCT
cn["A20"] = "Panama total per transit"; cn["A20"].font = BOLD
cn["C20"] = "=SUM(C7:C14)"; cn["D20"] = "=SUM(D7:D14)"
for c in ("C20", "D20"):
    cn[c].number_format = USD; cn[c].font = BOLD; cn[c].fill = OUT_FILL

cn["A24"] = "SUEZ CANAL — petroleum-product-tanker bands (SDR per SCNT), Circular 7/2023"; cn["A24"].font = H2
header(cn, 25, ["Band", "Band size (SCNT)", "Laden SDR/t", "Ballast SDR/t", "SCNT in band", "Note"])
bands = ["First 5,000", "Next 5,000", "Next 10,000", "Next 20,000", "Next 30,000", "Next 50,000", "The rest"]
for i, (lab, (size, l, b)) in enumerate(zip(bands, AN.SUEZ_BANDS)):
    rr = 26 + i
    cn.cell(row=rr, column=1, value=lab)
    c = cn.cell(row=rr, column=2, value=size if size < 10**8 else 10**9); c.font = BLUE_IN; c.number_format = NUM
    for col, v in ((3, l), (4, b)):
        c = cn.cell(row=rr, column=col, value=v); c.font = BLUE_IN; c.number_format = NUM2
    cum = f"SUM($B$26:B{rr-1})" if i else "0"
    cn.cell(row=rr, column=5, value=f"=MAX(0,MIN(B{rr},SCNT-{cum}))").number_format = NUM
cn["F26"] = "Source: SCA Circular 7/2023 (eff. 15-Jan-2024) via shipping-agent mirror; see data/README.md"
cn["A35"] = "Base dues (SDR)"; cn["C35"] = "=SUMPRODUCT(C26:C32,$E$26:$E$32)"; cn["D35"] = "=SUMPRODUCT(D26:D32,$E$26:$E$32)"
cn["A36"] = "Tanker surcharge"; cn["C36"] = "=SuezSurL"; cn["D36"] = "=SuezSurB"
cn["A37"] = "Dues incl. surcharge (SDR)"; cn["C37"] = "=C35*(1+C36)"; cn["D37"] = "=D35*(1+D36)"
cn["A38"] = "USD per SDR"; cn["C38"] = "=SDR_USD"; cn["D38"] = "=SDR_USD"
cn["A39"] = "Transit dues ($)"; cn["C39"] = "=C37*C38"; cn["D39"] = "=D37*D38"
cn["A40"] = "Ancillaries: pilotage, tugs, mooring, agency ($, estimate)"; cn["C40"] = "=SuezAnc"; cn["D40"] = "=SuezAnc"
cn["A41"] = "War-risk additional premium ($, illustrative)"; cn["C41"] = "=WarAP*HullValue"; cn["D41"] = "=WarAP*HullValue"
cn["A45"] = "Suez total per transit"; cn["A45"].font = BOLD
cn["C45"] = "=SUM(C39:C41)"; cn["D45"] = "=SUM(D39:D41)"
for rr in range(35, 46):
    for col in "CD":
        c = cn[f"{col}{rr}"]
        if c.value is not None:
            c.number_format = PCT if rr == 36 else ('0.0000' if rr == 38 else (NUM if rr in (35, 37) else USD))
for c in ("C45", "D45"):
    cn[c].font = BOLD; cn[c].fill = OUT_FILL
cn["C34"], cn["D34"] = "Laden", "Ballast"
cn["C34"].font = BOLD; cn["D34"].font = BOLD
cn["F36"] = "Chemical-tanker surcharge per safety4sea reporting of the same circular; bands are the product-tanker schedule (closest published proxy)."
cn["F36"].alignment = Alignment(wrap_text=True)

# =============================================================================
# Engine (row-wise voyage calculator)
# =============================================================================
title(eng, "Calculation engine — one full round-voyage P&L per row",
      "Each row recomputes the whole voyage with its own lever values (cols C-K). Sensitivity and Breakeven sheets read from here. "
      "Row 7 = active scenario and must equal Voyage_PnL (check on that sheet).")
ecols = ["Case", "Group", "Route #", "Brent $/bbl", "Freight mult", "Port-time mult", "Util FH", "Util BH",
         "Speed kn", "Port-cost mult", "EUA EUR/t",
         "Distance/leg nm", "ECA share", "Sea days (both legs)", "Speed factor", "Backhaul laden?",
         "Sea fuel t", "VLSFO t", "MGO sea t", "Port ops per region", "Port days A", "Port days B",
         "Canal days", "MGO port t", "Bunker cost $", "Port cost $", "Canal cost $", "EU ETS cost $",
         "Qty FH t", "Qty BH t", "Rate FH $/t", "Rate BH $/t", "Gross freight $", "Commission $",
         "Voyage costs $", "Voyage days", "TCE $/day", "Voyage result $", "Break-even FH rate $/t"]
header(eng, 6, ecols)
eng.row_dimensions[6].height = 45
IDX = lambda col, rrow: f"INDEX(Routes!${col}$6:${col}$9,$C{rrow})"


def engine_row(rr, case, group, route="Sw_Route", brent="Brent", fm="FreightMult", pm="PortMult",
               ufh="UtilFH", ubh="UtilBH", spd="Speed", pcm="PortCostMult", eua="EUA"):
    vals = [case, group, f"={route}" if isinstance(route, str) else route]
    for v in (brent, fm, pm, ufh, ubh, spd, pcm, eua):
        vals.append(f"={v}" if isinstance(v, str) else v)
    f = {
        "L": f"={IDX('E', rr)}",
        "M": f"={IDX('F', rr)}/L{rr}",
        "N": f"=2*L{rr}*(1+SeaMargin)/(I{rr}*24)",
        "O": f"=(I{rr}/RefSpeed)^SpeedExp",
        "P": f"=IF(H{rr}>0,1,0)",
        "Q": f"=N{rr}/2*O{rr}*(ConsLaden+IF(P{rr}=1,ConsLaden,ConsBallast))",
        "R": f"=Q{rr}*(1-M{rr})",
        "S": f"=Q{rr}*M{rr}",
        "T": f"=1+P{rr}",
        "U": f"={IDX('L', rr)}*T{rr}*F{rr}",
        "V": f"={IDX('M', rr)}*T{rr}*F{rr}",
        "W": f"={IDX('H', rr)}*2",
        "X": f"=(U{rr}+V{rr}+W{rr})*ConsPort",
        "Y": f"=R{rr}*D{rr}*VLSFO_Mult+(S{rr}+X{rr})*D{rr}*MGO_Mult",
        "AN": f"=LOA_ft*Port_Tariffs!$C$9*" + dock_formula(f"({IDX('L', rr)}*{IDX('X', rr)}*F{rr})"),
        "AO": f"=LOA_ft*Port_Tariffs!$C$9*" + dock_formula(f"({IDX('M', rr)}*{IDX('Y', rr)}*F{rr})"),
        "Z": (f"=({IDX('N', rr)}+AN{rr}+MIN(AC{rr}*{IDX('T', rr)},{IDX('V', rr)})"
              f"+{IDX('O', rr)}+AO{rr}+MIN(AC{rr}*{IDX('U', rr)},{IDX('W', rr)})"
              f"+IF(P{rr}=1,{IDX('N', rr)}+AN{rr}+MIN(AD{rr}*{IDX('T', rr)},{IDX('V', rr)})"
              f"+{IDX('O', rr)}+AO{rr}+MIN(AD{rr}*{IDX('U', rr)},{IDX('W', rr)}),0))*J{rr}"),
        "AA": f"={IDX('R', rr)}+IF(P{rr}=1,{IDX('R', rr)},{IDX('S', rr)})",
        "AB": f"=((R{rr}*CO2_VLSFO+S{rr}*CO2_MGO)*{IDX('I', rr)}+(U{rr}*{IDX('P', rr)}+V{rr}*{IDX('Q', rr)})*ConsPort*CO2_MGO)*PhaseIn*K{rr}*FX",
        "AC": f"=Capacity*G{rr}",
        "AD": f"=Capacity*H{rr}",
        "AE": f"={IDX('J', rr)}*E{rr}",
        "AF": f"={IDX('K', rr)}*E{rr}",
        "AG": f"=AC{rr}*AE{rr}+AD{rr}*AF{rr}",
        "AH": f"=AG{rr}*Commission",
        "AI": f"=Y{rr}+Z{rr}+AA{rr}+AB{rr}+MiscVoy",
        "AJ": f"=N{rr}+U{rr}+V{rr}+W{rr}",
        "AK": f"=(AG{rr}-AH{rr}-AI{rr})/AJ{rr}",
        "AL": f"=AG{rr}-AH{rr}-AI{rr}-Benchmark*AJ{rr}",
        "AM": f"=IF(AC{rr}>0,((Benchmark*AJ{rr}+AI{rr})/(1-Commission)-AD{rr}*AF{rr})/AC{rr},0)",
    }
    for j, v in enumerate(vals):
        c = eng.cell(row=rr, column=1 + j, value=v)
        if j >= 2:
            c.font = GREEN if isinstance(v, str) else BLUE_IN
    fm_map = {"D": NUM2, "E": MULT, "F": MULT, "G": PCT, "H": PCT, "I": NUM1, "J": MULT, "K": NUM1,
              "L": NUM, "M": PCT, "N": NUM1, "O": '0.000', "Q": NUM, "R": NUM, "S": NUM, "U": NUM1, "V": NUM1,
              "W": NUM1, "X": NUM, "Y": USD, "Z": USD, "AA": USD, "AB": USD, "AC": NUM, "AD": NUM, "AE": USD2,
              "AF": USD2, "AG": USD, "AH": USD, "AI": USD, "AJ": NUM1, "AK": USD, "AL": USD, "AM": USD2,
              "AN": USD, "AO": USD}
    for col, formula in f.items():
        eng[f"{col}{rr}"] = formula
    for col, fmt in fm_map.items():
        eng[f"{col}{rr}"].number_format = fmt
    eng[f"AK{rr}"].fill = OUT_FILL


ER = {}
rr = 7
engine_row(rr, "Active scenario", "base"); ER["base"] = rr
# tornado rows: (label, lever key in engine_row, low expr, high expr, unit, display fmt)
TOR = [
    ("Freight rate (market)", "fm", "Freight_Low", "Freight_High", "x base rate", MULT),
    ("Bunker price (Brent)", "brent", "Brent_P10", "Brent_P90", "$/bbl", USD2),
    ("Port time (congestion)", "pm", 0.75, 1.75, "x base port days", MULT),
    ("Fronthaul utilization", "ufh", 0.70, 0.95, "% capacity", PCT),
    ("Backhaul utilization", "ubh", 0.25, 0.75, "% capacity", PCT),
    ("Service speed", "spd", 12.0, 14.5, "knots", NUM1),
    ("Port disbursements", "pcm", 0.75, 1.25, "x base", MULT),
    ("EU carbon price (EUA)", "eua", 50.0, 110.0, "EUR/t", NUM1),
]
rr = 9
eng.cell(row=rr - 1, column=1, value="TORNADO CASES (one lever flexed, all others at active values)").font = BOLD
for lab, key, lo, hi, unit, fmt in TOR:
    engine_row(rr, f"{lab} — low", "tornado", **{key: lo}); ER[(lab, "lo")] = rr; rr += 1
    engine_row(rr, f"{lab} — high", "tornado", **{key: hi}); ER[(lab, "hi")] = rr; rr += 1
rr += 1
eng.cell(row=rr, column=1, value="ELASTICITY CASES (+10% on one lever)").font = BOLD; rr += 1
ELA = [("Freight rate", "fm", "FreightMult"), ("Bunker price", "brent", "Brent"), ("Port time", "pm", "PortMult"),
       ("Fronthaul utilization", "ufh", "UtilFH"), ("Backhaul utilization", "ubh", "UtilBH"),
       ("Service speed", "spd", "Speed"), ("Port disbursements", "pcm", "PortCostMult")]
for lab, key, base_nm in ELA:
    engine_row(rr, f"{lab} +10%", "elasticity", **{key: f"{base_nm}*1.1"}); ER[("ela", lab)] = rr; rr += 1
rr += 1
eng.cell(row=rr, column=1, value="ROUTE COMPARISON (active levers, each route)").font = BOLD; rr += 1
for rid in AN.ROUTES:
    engine_row(rr, f"Route {rid}", "routes", route=rid); ER[("route", rid)] = rr; rr += 1
rr += 1
BE_BRENTS = list(range(50, 135, 10))
BE_PORT = [1.0, 1.5, 2.0]
BE_UTIL = [0.70, 0.85, 0.95]
eng.cell(row=rr, column=1, value="BREAK-EVEN CASES (Brent x port-time; Brent x fronthaul utilization)").font = BOLD; rr += 1
for pmv in BE_PORT:
    for b in BE_BRENTS:
        engine_row(rr, f"BE port {pmv}x Brent {b}", "be_port", brent=b, pm=pmv); ER[("bep", pmv, b)] = rr; rr += 1
for u in BE_UTIL:
    for b in BE_BRENTS:
        engine_row(rr, f"BE util {u:.0%} Brent {b}", "be_util", brent=b, ufh=u); ER[("beu", u, b)] = rr; rr += 1
for rid in AN.ROUTES:
    for b in BE_BRENTS:
        engine_row(rr, f"BE route {rid} Brent {b}", "be_route", route=rid, brent=b); ER[("ber", rid, b)] = rr; rr += 1
rr += 1
GRID_F = [0.8, 0.9, 1.0, 1.1, 1.2, 1.3]
GRID_B = [60, 70, 80, 90, 100, 110, 120]
eng.cell(row=rr, column=1, value="TCE GRID CASES (freight multiplier x Brent)").font = BOLD; rr += 1
for fmv in GRID_F:
    for b in GRID_B:
        engine_row(rr, f"Grid {fmv}x / ${b}", "grid", fm=fmv, brent=b); ER[("grid", fmv, b)] = rr; rr += 1
rr += 1
SPEEDS = [10.0, 10.5, 11.0, 11.5, 12.0, 12.5, 13.0, 13.5, 14.0, 14.5, 15.0]
SP_B = [65, 82, 104]
eng.cell(row=rr, column=1, value="SPEED CASES (speed x Brent)").font = BOLD; rr += 1
for b in SP_B:
    for s in SPEEDS:
        engine_row(rr, f"Speed {s} / ${b}", "speed", spd=s, brent=b); ER[("spd", b, s)] = rr; rr += 1
eng.freeze_panes = "C7"
eng.column_dimensions["A"].width = 32; eng.column_dimensions["B"].width = 10
for c in range(3, len(ecols) + 1):
    eng.column_dimensions[L(c)].width = 12

# =============================================================================
# Voyage P&L (detailed, independent formula build-up for the active scenario)
# =============================================================================
title(pnl, "Voyage P&L — Time Charter Equivalent build-up (active scenario)",
      "Change scenario switches on the Inputs sheet. All values are formulas.")
pnl.column_dimensions["A"].width = 46; pnl.column_dimensions["B"].width = 18
pnl.column_dimensions["C"].width = 18; pnl.column_dimensions["D"].width = 58
RS = "Sw_Route"
RI = lambda col: f"INDEX(Routes!${col}$6:${col}$9,{RS})"
_dockA = f"LOA_ft*Port_Tariffs!$C$9*" + dock_formula(f"({RI('L')}*{RI('X')}*PortMult)")
_dockB = f"LOA_ft*Port_Tariffs!$C$9*" + dock_formula(f"({RI('M')}*{RI('Y')}*PortMult)")
_opA = lambda q: f"{RI('N')}+{_dockA}+MIN({q}*{RI('T')},{RI('V')})"
_opB = lambda q: f"{RI('O')}+{_dockB}+MIN({q}*{RI('U')},{RI('W')})"
PORT_FORMULA = (f"=({_opA('B23')}+{_opB('B23')}+IF(UtilBH>0,{_opA('B24')}+{_opB('B24')},0))*PortCostMult")
lines = [
    ("SCENARIO", None, None, None, "h"),                                                      # 7
    ("Route", f"=INDEX(Routes!$B$6:$B$9,{RS})", None, "", "t"),                               # 8
    ("Brent / VLSFO / MGO", "=Brent", '=TEXT(Brent*VLSFO_Mult,"$#,##0")&" / "&TEXT(Brent*MGO_Mult,"$#,##0")&" per t"', "$/bbl; bunker $/t derived", NUM2),  # 9
    ("Freight multiplier / port-time multiplier", "=FreightMult", "=PortMult", "", MULT),      # 10
    ("Utilization fronthaul / backhaul", "=UtilFH", "=UtilBH", "", PCT),                       # 11
    ("", None, None, None, "b"),                                                               # 12
    ("VOYAGE TIME", "Days", "", "", "h"),                                                      # 13
    ("Distance per leg (nm)", f"={RI('E')}", None, "Routes sheet", NUM),                       # 14
    ("Sea days, fronthaul", "=B14*(1+SeaMargin)/(Speed*24)", None, "distance x (1 + sea margin) / (speed x 24)", NUM1),   # 15
    ("Sea days, backhaul", "=B15", None, "same distance", NUM1),                               # 16
    ("Port days, region A", f"={RI('L')}*(1+IF(UtilBH>0,1,0))*PortMult", None, "days per operation x operations x port-time multiplier", NUM1),  # 17
    ("Port days, region B", f"={RI('M')}*(1+IF(UtilBH>0,1,0))*PortMult", None, "an operation = a load or a discharge; a backhaul cargo adds one at each end", NUM1),  # 18
    ("Canal days", f"={RI('H')}*2", None, "two transits", NUM1),                               # 19
    ("Total voyage days", "=SUM(B15:B19)", None, "", NUM1),                                    # 20
    ("", None, None, None, "b"),                                                               # 21
    ("REVENUE", "$", "", "", "h"),                                                             # 22
    ("Fronthaul cargo (t) x rate ($/t)", "=Capacity*UtilFH", f"={RI('J')}*FreightMult", "B = tonnes, C = $/t", NUM),   # 23
    ("Backhaul cargo (t) x rate ($/t)", "=Capacity*UtilBH", f"={RI('K')}*FreightMult", "", NUM),                        # 24
    ("Gross freight", "=B23*C23+B24*C24", None, "", USD),                                      # 25
    ("Less: commissions", "=-B25*Commission", None, "address + brokerage", USD),               # 26
    ("Net freight", "=B25+B26", None, "", USD),                                                # 27
    ("", None, None, None, "b"),                                                               # 28
    ("VOYAGE COSTS", "$", "t", "", "h"),                                                       # 29
    ("Speed factor (consumption vs reference)", "=(Speed/RefSpeed)^SpeedExp", None, "cube law", '0.000'),  # 30
    ("Sea fuel, both legs (t)", "=B15*B30*ConsLaden+B16*B30*IF(UtilBH>0,ConsLaden,ConsBallast)", None, "", NUM),  # 31
    ("ECA share of sea miles", f"={RI('F')}/B14", None, "burns MGO", PCT),                     # 32
    ("VLSFO cost (sea, outside ECA)", "=C33*Brent*VLSFO_Mult", "=B31*(1-B32)", "C = tonnes", USD),  # 33
    ("MGO cost (ECA + port + canal)", "=C34*Brent*MGO_Mult", "=B31*B32+(B17+B18+B19)*ConsPort", "C = tonnes", USD),  # 34
    ("Port disbursements", PORT_FORMULA, None, "Published Port Houston / Rotterdam tariffs + estimated ancillaries (Port_Tariffs sheet)", USD),  # 35
    ("Canal tolls & transit costs", f"={RI('R')}+IF(UtilBH>0,{RI('R')},{RI('S')})", None, "Canal_Tolls sheet", USD),  # 36
    ("EU ETS allowances", f"=((C33*CO2_VLSFO+B31*B32*CO2_MGO)*{RI('I')}+(B17*{RI('P')}+B18*{RI('Q')})*ConsPort*CO2_MGO)*PhaseIn*EUA*FX", None, "50% of extra-EU voyage emissions + 100% at EU berth", USD),  # 37
    ("Misc. voyage costs", "=MiscVoy", None, "", USD),                                         # 38
    ("Total voyage costs", "=SUM(B33:B38)", None, "", USD),                                    # 39
    ("", None, None, None, "b"),                                                               # 40
    ("RESULT", "", "", "", "h"),                                                               # 41
    ("Voyage net revenue (net freight - voyage costs)", "=B27-B39", None, "", USD),            # 42
    ("TIME CHARTER EQUIVALENT ($/day)", "=B42/B20", None, "the headline commercial metric", USD),  # 43
    ("Daily vessel cost benchmark ($/day)", "=Benchmark", None, "OPEX + capital, or T/C-in hire", USD),  # 44
    ("TCE margin over benchmark ($/day)", "=B43-B44", None, "", USD),                          # 45
    ("Voyage result after vessel cost ($)", "=B45*B20", None, "", USD),                        # 46
    ("Break-even fronthaul freight rate ($/t)", "=((B44*B20+B39)/(1-Commission)-B24*C24)/B23", None, "rate at which TCE = benchmark", USD2),  # 47
    ("Headroom: base fronthaul rate vs break-even", "=C23/B47-1", None, "how far the rate can fall before the voyage stops covering vessel cost", PCT),  # 48
    ("", None, None, None, "b"),                                                               # 49
    ("COST STRUCTURE (% of total voyage costs)", "", "", "", "h"),                             # 50
    ("Bunkers", "=(B33+B34)/B39", None, "", PCT),                                              # 51
    ("Port", "=B35/B39", None, "", PCT),                                                       # 52
    ("Canal", "=B36/B39", None, "", PCT),                                                      # 53
    ("EU ETS", "=B37/B39", None, "", PCT),                                                     # 54
    ("Misc", "=B38/B39", None, "", PCT),                                                       # 55
    ("Voyage costs as % of net freight", "=B39/B27", None, "", PCT),                           # 56
    ("", None, None, None, "b"),                                                               # 57
    ("MODEL CHECK", "", "", "", "h"),                                                          # 58
    ("Engine TCE (active-scenario row)", f"=Engine!AK{ER['base']}", None, "", USD),            # 59
    ("Difference (must be 0)", "=ROUND(B43-B59,6)", '=IF(B60=0,"OK","CHECK")', "", USD2),      # 60
]
for i, (lab, b, c, d, kind) in enumerate(lines):
    rrow = 7 + i
    if kind == "h":
        cell = pnl.cell(row=rrow, column=1, value=lab); cell.font = HDR_FONT
        for col in range(1, 5):
            pnl.cell(row=rrow, column=col).fill = HDR_FILL
        if b:
            pnl.cell(row=rrow, column=2, value=b).font = HDR_FONT
        if c:
            pnl.cell(row=rrow, column=3, value=c).font = HDR_FONT
        continue
    if kind == "b":
        continue
    pnl.cell(row=rrow, column=1, value=lab)
    if b is not None:
        cb = pnl.cell(row=rrow, column=2, value=b)
        cb.number_format = kind if kind != "t" else "General"
    if c is not None:
        cc = pnl.cell(row=rrow, column=3, value=c)
        cc.number_format = kind if kind != "t" else "General"
    pnl.cell(row=rrow, column=4, value=d).font = Font(name=FONT, size=9, color="595959")
# fix a few formats
pnl["C23"].number_format = USD2; pnl["C24"].number_format = USD2
pnl["C33"].number_format = NUM; pnl["C34"].number_format = NUM
pnl["C9"].number_format = "General"; pnl["C11"].number_format = PCT
for rrow in (25, 27, 39, 42, 46):
    pnl[f"A{rrow}"].font = BOLD; pnl[f"B{rrow}"].font = BOLD
for rrow in (43, 47):
    pnl[f"A{rrow}"].font = Font(name=FONT, bold=True, size=12, color=NAVY)
    pnl[f"B{rrow}"].font = Font(name=FONT, bold=True, size=12, color=NAVY)
    pnl[f"B{rrow}"].fill = KEY_FILL
pnl.conditional_formatting.add("B46", CellIsRule(operator="lessThan", formula=["0"], fill=WARN_FILL))
pnl.conditional_formatting.add("C60", CellIsRule(operator="equal", formula=['"CHECK"'], fill=WARN_FILL))
pnl.freeze_panes = "A7"

# Route comparison block on P&L sheet
pnl["F6"] = "ROUTE COMPARISON — active levers applied to every route"; pnl["F6"].font = H2
rc_hdr = ["Route", "Days", "Gross freight", "Bunkers", "Port", "Canal", "EU ETS", "Misc", "TCE $/day",
          "Result $", "Break-even FH $/t"]
header(pnl, 7, rc_hdr, start=6)
src_cols = ["AG", "Y", "Z", "AA", "AB", None, "AK", "AL", "AM"]
for i, rid in enumerate(AN.ROUTES):
    rrow = 8 + i
    er = ER[("route", rid)]
    pnl.cell(row=rrow, column=6, value=f"=INDEX(Routes!$B$6:$B$9,{rid})")
    pnl.cell(row=rrow, column=7, value=f"=Engine!AJ{er}").number_format = NUM1
    for j, col in enumerate(src_cols):
        v = f"=Engine!{col}{er}" if col else "=MiscVoy"
        c = pnl.cell(row=rrow, column=8 + j, value=v)
        c.number_format = USD2 if col == "AM" else USD
        c.font = GREEN
for col, w in zip("FGHIJKLMNOP", (26, 8, 13, 12, 11, 11, 11, 10, 11, 12, 12)):
    pnl.column_dimensions[col].width = w
bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100
bc.title = "Voyage cost build-up by route ($)"; bc.height = 8; bc.width = 18
bc.add_data(Reference(pnl, min_col=9, max_col=13, min_row=7, max_row=11), titles_from_data=True)
bc.set_categories(Reference(pnl, min_col=6, min_row=8, max_row=11))
pnl.add_chart(bc, "F14")
bc2 = BarChart(); bc2.type = "bar"; bc2.title = "TCE by route ($/day)"; bc2.height = 7; bc2.width = 18
bc2.add_data(Reference(pnl, min_col=14, min_row=7, max_row=11), titles_from_data=True)
bc2.set_categories(Reference(pnl, min_col=6, min_row=8, max_row=11)); bc2.legend = None
pnl.add_chart(bc2, "F31")

# =============================================================================
# Sensitivity
# =============================================================================
title(sens, "Sensitivity — what drives TCE?",
      "Tornado: each lever moved across a plausible range with everything else at the active scenario. "
      "Freight and Brent ranges are data-derived (Data_PPI, Data_Brent); the rest are judgment bands.")
sens.column_dimensions["A"].width = 26
for col in "BCDEFGHI":
    sens.column_dimensions[col].width = 13
sens["A5"] = "Base TCE (active scenario)"; sens["B5"] = f"=Engine!AK{ER['base']}"; sens["B5"].number_format = USD
sens["A5"].font = BOLD; sens["B5"].font = BOLD; sens["B5"].fill = KEY_FILL
sens["A7"] = "A. Tornado input table (unsorted)"; sens["A7"].font = H2
header(sens, 8, ["Lever", "Low value", "High value", "Unit", "TCE at low", "TCE at high", "Swing $/day", "Sort key"])
for i, (lab, key, lo, hi, unit, fmt) in enumerate(TOR):
    rrow = 9 + i
    lo_r, hi_r = ER[(lab, "lo")], ER[(lab, "hi")]
    colmap = {"fm": "E", "brent": "D", "pm": "F", "ufh": "G", "ubh": "H", "spd": "I", "pcm": "J", "eua": "K"}[key]
    sens.cell(row=rrow, column=1, value=lab)
    sens.cell(row=rrow, column=2, value=f"=Engine!{colmap}{lo_r}").number_format = fmt
    sens.cell(row=rrow, column=3, value=f"=Engine!{colmap}{hi_r}").number_format = fmt
    sens.cell(row=rrow, column=4, value=unit)
    sens.cell(row=rrow, column=5, value=f"=Engine!AK{lo_r}").number_format = USD
    sens.cell(row=rrow, column=6, value=f"=Engine!AK{hi_r}").number_format = USD
    sens.cell(row=rrow, column=7, value=f"=ABS(F{rrow}-E{rrow})").number_format = USD
    sens.cell(row=rrow, column=8, value=f"=G{rrow}+ROW()/1000000").number_format = '0.000'
    for col in (2, 3, 5, 6):
        sens.cell(row=rrow, column=col).font = GREEN
n_t = len(TOR)
sens["A19"] = "B. Tornado — ranked by swing (chart source)"; sens["A19"].font = H2
header(sens, 20, ["Lever", "Low -> TCE change", "High -> TCE change", "Swing $/day", "Low value", "High value", "Rank"])
for k in range(n_t):
    rrow = 21 + k
    m = f"MATCH(LARGE($H$9:$H${8+n_t},{k+1}),$H$9:$H${8+n_t},0)"
    sens.cell(row=rrow, column=1, value=f"=INDEX($A$9:$A${8+n_t},{m})")
    sens.cell(row=rrow, column=2, value=f"=INDEX($E$9:$E${8+n_t},{m})-$B$5").number_format = USD
    sens.cell(row=rrow, column=3, value=f"=INDEX($F$9:$F${8+n_t},{m})-$B$5").number_format = USD
    sens.cell(row=rrow, column=4, value=f"=INDEX($G$9:$G${8+n_t},{m})").number_format = USD
    sens.cell(row=rrow, column=5, value=f"=INDEX($B$9:$B${8+n_t},{m})").number_format = "General"
    sens.cell(row=rrow, column=6, value=f"=INDEX($C$9:$C${8+n_t},{m})").number_format = "General"
    sens.cell(row=rrow, column=7, value=k + 1)
    for col in (5, 6):
        sens.cell(row=rrow, column=col).number_format = '#,##0.00'
tc = BarChart(); tc.type = "bar"; tc.grouping = "clustered"; tc.overlap = 100
tc.title = "Tornado: change in TCE vs base ($/day)"; tc.height = 9; tc.width = 20
tc.add_data(Reference(sens, min_col=2, max_col=3, min_row=20, max_row=20 + n_t), titles_from_data=True)
tc.set_categories(Reference(sens, min_col=1, min_row=21, max_row=20 + n_t))
tc.y_axis.number_format = '#,##0'; tc.x_axis.scaling.orientation = "maxMin"
tc.y_axis.majorGridlines = None
tc.x_axis.delete = False; tc.y_axis.delete = False
sens.add_chart(tc, "J5")

sens["A32"] = "C. Leverage — TCE change for a +10% move in each lever (apples-to-apples)"; sens["A32"].font = H2
header(sens, 33, ["Lever", "Base value", "+10% value", "TCE change $/day", "Per 1% move $/day"])
for i, (lab, key, base_nm) in enumerate(ELA):
    rrow = 34 + i
    er = ER[("ela", lab)]
    colmap = {"fm": "E", "brent": "D", "pm": "F", "ufh": "G", "ubh": "H", "spd": "I", "pcm": "J"}[key]
    sens.cell(row=rrow, column=1, value=lab)
    sens.cell(row=rrow, column=2, value=f"={base_nm}").number_format = '#,##0.00'
    sens.cell(row=rrow, column=3, value=f"=Engine!{colmap}{er}").number_format = '#,##0.00'
    sens.cell(row=rrow, column=4, value=f"=Engine!AK{er}-$B$5").number_format = USD
    sens.cell(row=rrow, column=5, value=f"=D{rrow}/10").number_format = USD
ec = BarChart(); ec.type = "bar"; ec.title = "TCE change for +10% in each lever ($/day)"
ec.height = 8; ec.width = 20; ec.legend = None
ec.add_data(Reference(sens, min_col=4, min_row=33, max_row=33 + len(ELA)), titles_from_data=True)
ec.set_categories(Reference(sens, min_col=1, min_row=34, max_row=33 + len(ELA)))
ec.x_axis.scaling.orientation = "maxMin"; ec.x_axis.delete = False; ec.y_axis.delete = False
sens.add_chart(ec, "J25")
sens["A43"] = ("Reading the two views together: the tornado answers 'what moved results most over realistic ranges'; "
               "the +10% table answers 'where does one unit of effort buy the most TCE'. Utilization and speed are "
               "partly controllable; freight and bunker prices are mostly not.")
sens["A43"].alignment = Alignment(wrap_text=True, vertical="top"); sens.merge_cells("A43:H47")

# =============================================================================
# Breakeven
# =============================================================================
title(be, "Break-even fronthaul freight rate ($/t) — the rate at which TCE equals the daily vessel cost benchmark",
      "Curves hold the active route and levers constant except the ones named. Linear in Brent because bunker cost is linear in price.")
be.column_dimensions["A"].width = 22
for c in range(2, 16):
    be.column_dimensions[L(c)].width = 12
be["A5"] = "A. By port-time multiplier (active route)"; be["A5"].font = H2
header(be, 6, ["Brent $/bbl"] + [f"{p:.1f}x port time" for p in BE_PORT])
for i, b in enumerate(BE_BRENTS):
    rrow = 7 + i
    be.cell(row=rrow, column=1, value=b).font = BLUE_IN
    for j, pmv in enumerate(BE_PORT):
        be.cell(row=rrow, column=2 + j, value=f"=Engine!AM{ER[('bep', pmv, b)]}").number_format = USD2
nb = len(BE_BRENTS)
lc = LineChart(); lc.title = "Break-even FH rate vs Brent, by port time"; lc.height = 8; lc.width = 16
lc.add_data(Reference(be, min_col=2, max_col=4, min_row=6, max_row=6 + nb), titles_from_data=True)
lc.set_categories(Reference(be, min_col=1, min_row=7, max_row=6 + nb))
lc.y_axis.title = "$/t"; lc.x_axis.title = "Brent $/bbl"; lc.x_axis.delete = False; lc.y_axis.delete = False
be.add_chart(lc, "G5")

be["A19"] = "B. By fronthaul utilization (active route)"; be["A19"].font = H2
header(be, 20, ["Brent $/bbl"] + [f"{u:.0%} utilization" for u in BE_UTIL])
for i, b in enumerate(BE_BRENTS):
    rrow = 21 + i
    be.cell(row=rrow, column=1, value=b).font = BLUE_IN
    for j, u in enumerate(BE_UTIL):
        be.cell(row=rrow, column=2 + j, value=f"=Engine!AM{ER[('beu', u, b)]}").number_format = USD2
lc = LineChart(); lc.title = "Break-even FH rate vs Brent, by utilization"; lc.height = 8; lc.width = 16
lc.add_data(Reference(be, min_col=2, max_col=4, min_row=20, max_row=20 + nb), titles_from_data=True)
lc.set_categories(Reference(be, min_col=1, min_row=21, max_row=20 + nb))
lc.y_axis.title = "$/t"; lc.x_axis.title = "Brent $/bbl"; lc.x_axis.delete = False; lc.y_axis.delete = False
be.add_chart(lc, "G22")

be["A33"] = "C. By route (active levers)"; be["A33"].font = H2
header(be, 34, ["Brent $/bbl"] + [f"=Routes!B{5+rid}" for rid in AN.ROUTES] + [f"Base rate R{rid}" for rid in AN.ROUTES])
for i, b in enumerate(BE_BRENTS):
    rrow = 35 + i
    be.cell(row=rrow, column=1, value=b).font = BLUE_IN
    for j, rid in enumerate(AN.ROUTES):
        be.cell(row=rrow, column=2 + j, value=f"=Engine!AM{ER[('ber', rid, b)]}").number_format = USD2
        be.cell(row=rrow, column=6 + j, value=f"=Routes!$J${5+rid}*FreightMult").number_format = USD2
lc = LineChart(); lc.title = "Break-even FH rate vs Brent, by route"; lc.height = 8; lc.width = 16
lc.add_data(Reference(be, min_col=2, max_col=5, min_row=34, max_row=34 + nb), titles_from_data=True)
lc.set_categories(Reference(be, min_col=1, min_row=35, max_row=34 + nb))
lc.y_axis.title = "$/t"; lc.x_axis.title = "Brent $/bbl"; lc.x_axis.delete = False; lc.y_axis.delete = False
be.add_chart(lc, "G39")

be["A47"] = "D. TCE grid ($/day): freight multiplier (rows) x Brent (columns), active route"; be["A47"].font = H2
header(be, 48, ["Freight mult \\ Brent"] + [f"${b}" for b in GRID_B])
for i, fmv in enumerate(GRID_F):
    rrow = 49 + i
    c = be.cell(row=rrow, column=1, value=fmv); c.font = BLUE_IN; c.number_format = MULT
    for j, b in enumerate(GRID_B):
        be.cell(row=rrow, column=2 + j, value=f"=Engine!AK{ER[('grid', fmv, b)]}").number_format = USD
gr = f"B49:{L(1+len(GRID_B))}{48+len(GRID_F)}"
be.conditional_formatting.add(gr, ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num",
                                                 mid_value=A["daily_cost_benchmark"], mid_color="FFFFFF",
                                                 end_type="max", end_color="63BE7B"))
be["A56"] = "Colour midpoint = $18,000/day benchmark (white). Red = below vessel cost; green = above."
be["A56"].font = Font(name=FONT, italic=True, size=9)

be["A59"] = "E. Speed optimization: TCE ($/day) by service speed and Brent (active route)"; be["A59"].font = H2
header(be, 60, ["Speed (kn)"] + [f"Brent ${b}" for b in SP_B])
for i, s in enumerate(SPEEDS):
    rrow = 61 + i
    be.cell(row=rrow, column=1, value=s).font = BLUE_IN
    for j, b in enumerate(SP_B):
        be.cell(row=rrow, column=2 + j, value=f"=Engine!AK{ER[('spd', b, s)]}").number_format = USD
rrow = 61 + len(SPEEDS)
be.cell(row=rrow, column=1, value="TCE-maximizing speed").font = BOLD
for j in range(len(SP_B)):
    col = L(2 + j)
    be.cell(row=rrow, column=2 + j,
            value=f"=INDEX($A$61:$A${60+len(SPEEDS)},MATCH(MAX({col}61:{col}{60+len(SPEEDS)}),{col}61:{col}{60+len(SPEEDS)},0))").number_format = NUM1
lc = LineChart(); lc.title = "TCE vs speed at three bunker prices"; lc.height = 8; lc.width = 16
lc.add_data(Reference(be, min_col=2, max_col=4, min_row=60, max_row=60 + len(SPEEDS)), titles_from_data=True)
lc.set_categories(Reference(be, min_col=1, min_row=61, max_row=60 + len(SPEEDS)))
lc.x_axis.title = "knots"; lc.y_axis.title = "$/day"; lc.x_axis.delete = False; lc.y_axis.delete = False
be.add_chart(lc, "G58")

# =============================================================================
# Cover
# =============================================================================
title(cover, "Chemical Tanker Voyage Economics Model")
cover.column_dimensions["A"].width = 3; cover.column_dimensions["B"].width = 44
cover.column_dimensions["C"].width = 22; cover.column_dimensions["D"].width = 70
cover["A1"] = None
cover["B1"] = "Chemical Tanker Voyage Economics Model"; cover["B1"].font = H1
cover["B2"] = "Round-voyage TCE P&L, sensitivity/tornado, and break-even freight analysis — Sean Burgess, Sep-2026"
cover["B2"].font = Font(name=FONT, size=11, color="404040")
cover["B3"] = DISCLAIMER; cover["B3"].font = Font(name=FONT, bold=True, color="C00000")
cover["B4"] = ("SIMULATED / ESTIMATED INPUTS: freight rates, port costs, port days, vessel particulars and fuel-price "
               "conversion factors are analyst assumptions, not observed data. Canal tolls come from published tariff "
               "extracts; Brent and the freight index are public series.")
cover["B4"].font = Font(name=FONT, bold=True, color="C00000"); cover["B4"].alignment = Alignment(wrap_text=True)
cover.merge_cells("B4:D5"); cover.row_dimensions[4].height = 30; cover.row_dimensions[5].height = 18
cover["B7"] = "HEADLINE — ACTIVE SCENARIO"; cover["B7"].font = H2
kpis = [("Route", "=Voyage_PnL!B8", "General"), ("Round-voyage days", "=Voyage_PnL!B20", NUM1),
        ("Gross freight", "=Voyage_PnL!B25", USD), ("Total voyage costs", "=Voyage_PnL!B39", USD),
        ("TCE ($/day)", "=Voyage_PnL!B43", USD), ("Vessel cost benchmark ($/day)", "=Benchmark", USD),
        ("Voyage result after vessel cost", "=Voyage_PnL!B46", USD),
        ("Break-even fronthaul rate ($/t)", "=Voyage_PnL!B47", USD2),
        ("Largest tornado driver", "=Sensitivity!A21", "General"),
        ("Model integrity check", "=Voyage_PnL!C60", "General")]
for i, (k, f, fmt) in enumerate(kpis):
    rrow = 8 + i
    cover.cell(row=rrow, column=2, value=k)
    c = cover.cell(row=rrow, column=3, value=f); c.number_format = fmt; c.font = Font(name=FONT, bold=True, color="008000")
    c.fill = OUT_FILL; c.border = BOX
cover["B20"] = "HOW TO USE"; cover["B20"].font = H2
howto = [
    "1. Go to Inputs and change the five yellow scenario switches (route, bunker case, freight case, congestion, utilization).",
    "2. Every assumption is a blue cell on Inputs / Routes / Canal_Tolls with its basis written next to it — overwrite any you disagree with.",
    "3. Voyage_PnL shows the full cost build-up and TCE for the active scenario, plus all four routes side by side.",
    "4. Sensitivity ranks the drivers (tornado) and shows leverage per +10% move; Breakeven shows break-even freight curves, a TCE heat-map and a speed curve.",
    "5. Engine is the row-wise calculator behind the charts: one full voyage P&L per row. Do not edit it unless extending the model.",
]
for i, t in enumerate(howto):
    cover.cell(row=21 + i, column=2, value=t)
cover["B27"] = "COLOUR KEY"; cover["B27"].font = H2
keys = [("Scenario switch", SWITCH_FILL, BLUE_IN), ("Hardcoded input / assumption", None, BLUE_IN),
        ("Key assumption", KEY_FILL, BLUE_IN), ("Formula", None, BLACK), ("Link from another sheet", None, GREEN),
        ("Output", OUT_FILL, BLACK)]
for i, (k, fill, font) in enumerate(keys):
    c = cover.cell(row=28 + i, column=2, value=k); c.font = font
    if fill:
        c.fill = fill
cover["B35"] = "DEFINITIONS"; cover["B35"].font = H2
defs = [
    ("TCE", "(Gross freight - commissions - voyage costs) / voyage days. Converts a voyage charter into a $/day figure comparable with time-charter hire."),
    ("Voyage costs", "Costs the shipowner bears on a voyage charter: bunkers, port disbursements, canal tolls, EU ETS allowances, misc."),
    ("Fronthaul / backhaul", "Round voyage A->B (main cargo) then B->A (return cargo at lower utilization and rate). Backhaul utilization 0% = ballast return."),
    ("Break-even rate", "Fronthaul $/t at which TCE equals the daily vessel-cost benchmark."),
    ("ECA", "Emission Control Area — 0.10% sulphur limit, so MGO is burned inside (North American ECA, North Sea/Channel)."),
]
for i, (k, d) in enumerate(defs):
    cover.cell(row=36 + i, column=2, value=k).font = BOLD
    c = cover.cell(row=36 + i, column=3, value=d); c.alignment = Alignment(wrap_text=True, vertical="top")
    cover.merge_cells(start_row=36 + i, start_column=3, end_row=36 + i, end_column=4)
    cover.row_dimensions[36 + i].height = 30
cover["B42"] = "SHEETS"; cover["B42"].font = H2
sheets = ["Inputs — switches and assumptions", "Routes — route & port-region tables", "Canal_Tolls — Panama & Suez per-transit build-up",
          "Voyage_PnL — detailed TCE build-up + route comparison", "Sensitivity — tornado + leverage",
          "Breakeven — break-even curves, TCE grid, speed curve", "Engine — row-wise calculator",
          "Data_Brent — FRED Brent monthly + statistics", "Data_PPI — BLS deep-sea freight PPI + percentile bands"]
for i, s in enumerate(sheets):
    cover.cell(row=43 + i, column=2, value=s)
cover.sheet_view.showGridLines = False

for ws in wb.worksheets:
    style_all(ws)
    ws.sheet_properties.tabColor = {"Cover": NAVY, "Inputs": "FFC000", "Voyage_PnL": "2E75B6",
                                    "Sensitivity": "2E75B6", "Breakeven": "2E75B6"}.get(ws.title, "A6A6A6")
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
import json
(HERE / "model" / "_engine_rows.json").write_text(json.dumps({str(k): v for k, v in ER.items()}))
print("saved", OUT, "engine rows", len(ER))
